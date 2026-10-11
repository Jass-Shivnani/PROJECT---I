"""Generate the complete Dione Capstone Report as .docx"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = 'Times New Roman'
doc = Document()

for s in doc.sections:
    s.top_margin = Cm(2.54)
    s.bottom_margin = Cm(2.54)
    s.left_margin = Cm(3.17)
    s.right_margin = Cm(2.54)

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(12)
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.5

def h1(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t.upper()); r.bold = True; r.font.size = Pt(20); r.font.name = FONT
    p.paragraph_format.space_before = Pt(24); p.paragraph_format.space_after = Pt(12)

def h2(t):
    p = doc.add_paragraph()
    r = p.add_run(t); r.bold = True; r.font.size = Pt(16); r.font.name = FONT
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(8)

def h3(t):
    p = doc.add_paragraph()
    r = p.add_run(t); r.bold = True; r.font.size = Pt(14); r.font.name = FONT
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)

def bd(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(t); r.font.size = Pt(12); r.font.name = FONT
    p.paragraph_format.line_spacing = 1.5

def ct(t, sz=12, b=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = b; r.font.size = Pt(sz); r.font.name = FONT

def ip(c):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'[Image Placeholder: {c}]')
    r.italic = True; r.font.size = Pt(11); r.font.name = FONT; r.font.color.rgb = RGBColor(128,128,128)

def pb(): doc.add_page_break()

def tbl(headers, rows):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers)); t.style = 'Table Grid'
    for j,h in enumerate(headers): t.cell(0,j).text = h
    for i,row in enumerate(rows):
        for j,val in enumerate(row): t.cell(i+1,j).text = val

# === TITLE PAGE ===
ct("VIVEKANAND EDUCATION SOCIETY'S INSTITUTE OF TECHNOLOGY", 14, True)
ct("(An Autonomous Institute Affiliated to University of Mumbai)", 12)
ct("Department of Computer Engineering", 14, True)
doc.add_paragraph(); ip("College Logo"); doc.add_paragraph()
ct("Project Report on", 14); ct("DIONE", 24, True)
ct("A Local Large Action Model (LAM) Orchestration Engine", 14)
doc.add_paragraph()
ct("In partial fulfillment of the Third Year (Semester-VI),", 12)
ct("Bachelor of Technology (B.Tech.) Degree in Computer Engineering", 12)
ct("at the University of Mumbai, Academic Year 2025-26", 12)
doc.add_paragraph()
ct("Project Mentor: Prof. Geocey Shejy", 12, True); doc.add_paragraph()
ct("Submitted by", 12, True)
for n in ["Riya Khialani D12A, 34","Jass Shivnani D12B, 50","Khushi Sadhuramani D12B, 44","Ruchika Dingria D12B, 70"]: ct(n)
ct("(2025-26)", 12, True); pb()

# === CERTIFICATE ===
h1("Certificate of Approval")
bd('This is to certify that Riya Khialani (D12A), Jass Shivnani (D12B), Khushi Sadhuramani (D12B), and Ruchika Dingria (D12B) of Third Year Computer Engineering studying under the University of Mumbai have satisfactorily presented the project on "Dione" as a part of the coursework of CAPSTONE PROJECT \u2013 I for Semester-VI under the guidance of Prof. Geocey Shejy in the year 2025-26.')
doc.add_paragraph(); ct("Date: _______________"); doc.add_paragraph()
tbl(["Internal Examiner","External Examiner"],[["___________________","___________________"]])
doc.add_paragraph()
tbl(["Project Mentor","Head of Department","Principal"],[["___________________","Dr. Mrs. Nupur Giri\n___________________","Dr. J. M. Nair\n___________________"]]); pb()

# === DECLARATION ===
h1("Declaration")
bd('We declare that this written submission represents our ideas in our own words and where others\u2019 ideas or words have been included, we have adequately cited and referenced the original sources. We also declare that we have adhered to all principles of academic honesty and integrity and have not misrepresented or fabricated or falsified any idea, data, fact, or source in our submission. We understand that any violation of the above will be cause for disciplinary action by the Institute.')
ct("Date: _______________"); pb()

# === ACKNOWLEDGEMENT ===
h1("Acknowledgement")
bd('We are thankful to our college Vivekanand Education Society\u2019s Institute of Technology for considering our project and extending help at all stages needed during our work of collecting information regarding the project.')
bd('It gives us immense pleasure to express our deep and sincere gratitude to Professor Geocey Shejy (Project Guide) for her kind help and valuable advice during the development of the project synopsis and for her guidance and suggestions.')
bd('We are deeply indebted to the Head of the Computer Department, Dr. (Mrs.) Nupur Giri, and our Principal, Dr. (Mrs.) J. M. Nair, for giving us this valuable opportunity to do this project.')
bd('We express our hearty thanks to them for their assistance without which it would have been difficult in finishing this project synopsis and project review successfully.')
bd('We convey our deep sense of gratitude to all teaching and non-teaching staff for their constant encouragement, support, and selfless help throughout the project work. It is a great pleasure to acknowledge the help and suggestions which we received from the Department of Computer Engineering.')
bd('We wish to express our profound thanks to all those who helped us in gathering information about the project. Our families too have provided moral support and encouragement several times.'); pb()

# === COURSE OUTCOMES ===
h1("Course Outcomes"); bd("Learners will be able to:")
tbl(["CO","Description"],[("CO1","Identify and define a real-world problem through systematic requirement analysis."),("CO2","Conduct a thorough literature review and benchmarking to propose innovative solutions."),("CO3","Develop a structured project plan with well-defined milestones and timelines."),("CO4","Design and implement a working prototype using appropriate engineering tools."),("CO5","Analyze experimental results, validate the solution, and optimize performance."),("CO6","Effectively document and present the project outcomes.")]); pb()

# === ABSTRACT ===
h1("Abstract")
bd('The rapid advancement of Large Language Models (LLMs) has significantly enhanced the ability of artificial intelligence systems to understand, generate, and reason over human language. However, these models remain fundamentally limited in their functionality as they operate primarily within isolated chat interfaces. While they can suggest solutions or generate outputs, they lack the capability to execute real-world digital actions, creating a critical limitation known as the Actuation Gap.')
bd('To address this limitation, this project introduces Dione, a Local Large Action Model (LAM) Orchestration Engine designed to enable secure and autonomous execution of tasks. Unlike existing cloud-based agentic AI systems that require sensitive user data to be transmitted to external servers, Dione operates entirely on the user\u2019s local device. It acts as an intelligent middleware that bridges the reasoning capabilities of local LLMs with system-level tools, APIs, and applications.')
bd('The proposed system emphasizes a privacy-first, modular, and secure architecture built on principles such as zero-trust local inference, sandboxed plugin execution, and complete data sovereignty. Through its micro-kernel design and ReAct-based operational loop, Dione ensures efficient task execution while maintaining strict control over user data.'); pb()

# === TABLE OF CONTENTS ===
h1("Table of Contents"); bd("[To be generated automatically in Word using Insert > Table of Contents.]"); pb()

# === CHAPTER 1 ===
h1("Chapter 1: Introduction")
bd('This chapter introduces the concept, background, and motivation behind the proposed system, Dione \u2013 a Local Large Action Model (LAM) Orchestration Engine.')
h2("1.1 Introduction to the Project")
bd('The rapid evolution of Artificial Intelligence has led to the development of Large Language Models capable of performing advanced reasoning, natural language understanding, and content generation. Despite these advancements, such models are inherently limited as they operate within isolated environments, primarily restricted to chat-based interactions. To overcome this limitation, Dione is designed as a Local LAM Orchestration Engine that bridges the gap between the reasoning capabilities of LLMs and the execution capabilities of system-level tools and applications.')
h2("1.2 Motivation for the Project")
bd('The primary motivation arises from the inherent limitation of current AI systems, which are capable of reasoning but incapable of acting. This creates a dependency on users to manually execute tasks even after receiving accurate guidance. Additionally, existing agentic AI systems rely heavily on cloud-based architectures, requiring users to share sensitive personal data with external servers. The motivation is to develop a local-first AI system that performs intelligent reasoning and executes tasks autonomously while ensuring complete data privacy.')
h2("1.3 Drawbacks of Existing Systems")
bd('Current AI-driven assistants suffer from several critical limitations. Most systems are cloud-dependent, requiring user data to be transmitted to remote servers, leading to exposure of sensitive information. These systems often require storage of high-privilege credentials on centralized platforms. The dependency on internet connectivity introduces latency and reduces reliability. Additionally, existing systems lack true autonomy for multi-step workflows.')
h2("1.4 Problem Definition")
bd('The problem addressed is the inability of current AI systems to bridge the gap between reasoning and execution while maintaining user privacy and system security. There is a need to design a system that can interpret natural language instructions, convert them into structured actions, and execute them directly within the user\u2019s environment, operating entirely on local hardware.')
h2("1.5 Relevance of the Project")
bd('Dione is highly relevant in the current technological landscape where there is a growing shift towards Agentic AI and Edge Computing. It enables autonomous task execution without exposing sensitive data to external entities, making it particularly relevant for personal productivity, enterprise automation, and privacy-sensitive sectors including finance, healthcare, and legal services.')
h2("1.6 Methodology Used")
bd('The development follows a modular approach based on a micro-kernel architecture. A ReAct (Reason + Act) loop is implemented for iterative reasoning and execution. The system integrates local LLMs through Ollama, uses RAG techniques with ChromaDB for context management, and handles task execution through a plugin-based architecture with Pydantic validation and sandboxing.'); pb()

# === CHAPTER 2 ===
h1("Chapter 2: Literature Survey")
bd('This chapter presents a detailed study of existing research, frameworks, and resources relevant to the development of Dione.')
h2("2.1 Research Papers")
h3("Paper 1: Toolformer (2023)")
bd('Abstract: Toolformer enables language models to learn how to use external tools such as APIs during inference. Inference: It enhances reasoning by integrating tools but lacks a complete execution pipeline and does not address privacy concerns.')
h3("Paper 2: ReAct (2023)")
bd('Abstract: ReAct combines reasoning and action generation in a loop. Inference: It provides a strong conceptual foundation for agentic AI but is only a prompting strategy and lacks system-level execution and security mechanisms.')
h3("Paper 3: Generative Agents (2023)")
bd('Abstract: Demonstrates agents capable of simulating human-like behavior using memory, planning, and reflection. Inference: Useful for Dione\u2019s Context Manager but simulation-focused and does not emphasize real-world task execution.')
h3("Paper 4: RAG (2020)")
bd('Abstract: RAG combines retrieval mechanisms with generative models to enhance accuracy. Inference: Critical for contextual memory in Dione but only enhances reasoning, not action execution.')
h2("2.2 Books / Articles Referred")
bd('Key references include "Attention Is All You Need" (Vaswani et al.), the GPT-4 Technical Report, articles on Agentic AI (2023-2024), and industry blogs from LangChain and Ollama. These collectively highlight the shift toward agentic AI but reveal dependency on cloud-based systems.')
h2("2.3 Interaction with Domain Experts")
bd('Interaction with faculty and professionals validated the feasibility of the proposed approach. Key insights: local LLM deployment is increasingly feasible; security layers are critical; plugin-based architectures are preferred for scalability; and there is strong industry demand for privacy-preserving AI systems.')
h2("2.4 Patent Search")
bd('Patent 1: Intelligent Personal Assistant System (2017) \u2013 Cloud-based, lacks local execution. Patent 2: Autonomous Task Execution Using AI Agents (2021) \u2013 Provides execution but lacks real-time reasoning. Patent 3: Secure Data Processing in AI Systems (2020) \u2013 Addresses security but not execution. Patent 4: Context-Aware Digital Assistant (2019) \u2013 Improves personalization but remains cloud-based.'); pb()

# === CHAPTER 3 ===
h1("Chapter 3: Requirement of Proposed System")
h2("3.1 Functional Requirements")
bd('The system shall accept natural language input, process it using a local LLM, convert outputs into structured JSON commands, select appropriate plugins, execute tasks (file retrieval, email sending, API access), support multi-step execution, provide feedback, maintain contextual memory via ChromaDB, and validate actions before execution.')
h2("3.2 Non-Functional Requirements")
bd('Performance: Time-to-Action < 3 seconds. Security: All actions validated. Privacy: No external data transmission. Reliability: Graceful error handling. Scalability: Easy plugin integration. Usability: Simple interface. Maintainability: Clean code structure.')
h2("3.3 Constraints")
bd('Hardware limitations (RAM/GPU dependent), model size constraints (quantization needed), CPU-only systems slower, security restrictions may limit automation, API rate limits, and some features require internet connectivity.')
h2("3.4 Hardware and Software Requirements")
bd('Hardware: Intel i5/Ryzen 5+, 16-32 GB RAM, NVIDIA RTX 3060/Apple Silicon, 50 GB SSD. Software: Python 3.10+, Ollama/Llama.cpp, ChromaDB, SQLite, PyTorch, Pydantic, Gmail API, Calendar API.')
h2("3.5 Techniques Utilized")
bd('ReAct Loop for iterative reasoning/execution, RAG for context awareness, Prompt Engineering for actionable outputs, Model Quantization for efficient inference, and Schema Validation using Pydantic.')
h2("3.6 Tools Utilized")
bd('Ollama for local LLMs, Llama.cpp for efficient execution, ChromaDB for vector memory, SQLite for logging, Python as core language, VS Code as IDE, Git/GitHub for version control.')
h2("3.7 SDGs Mapped")
bd('SDG 9: Innovation in local AI infrastructure. SDG 8: Productivity through task automation. SDG 16: Data privacy and security. SDG 4: Platform for learning advanced AI concepts.'); pb()

# === CHAPTER 4 ===
h1("Chapter 4: Proposed Design")
h2("4.1 Block Diagram"); ip("Figure 4.1 \u2013 Block Diagram of Dione System")
bd('The block diagram shows the high-level flow: User Interface accepts input, Context Manager enriches with memory, Local LLM generates JSON commands, Validation Layer checks security, Plugin Registry selects tools, and Execution Layer performs actions.')
h2("4.2 Modular Diagram"); ip("Figure 4.2 \u2013 Modular Diagram of Dione System")
bd('Modules include: Context Manager (ChromaDB memory), Model Adapter (Ollama/Llama.cpp), Plugin Registry (dynamic tool loading), Safety Guardrail (unsafe command detection), and Execution Engine (validated action execution).')
h2("4.3 System Design")
ip("Figure 4.3 \u2013 Use Case Diagram"); ip("Figure 4.4 \u2013 System Flowchart (ReAct Loop)")
bd('DFD Level 0: User \u2192 Dione \u2192 External Systems \u2192 Output. Level 1: Input \u2192 Context Manager \u2192 LLM \u2192 Validation \u2192 Plugin \u2192 Execution \u2192 Output.')
bd('The flowchart represents the ReAct loop: Start, Accept Input, Retrieve Context, Generate Action, Validate, Execute Plugin, Check Completion, Repeat if needed, Display Output, End.')
h2("4.4 Algorithms in Existing Systems")
bd('Intent Classification (Siri, Alexa), Sequence-to-Sequence Models, Chain-of-Thought Reasoning, ReAct Algorithm (used in Dione), and RAG for context enhancement. Existing systems focus on either reasoning or execution but fail to integrate both securely and locally.')
h2("4.5 Project Scheduling")
tbl(["Phase","Task","Duration"],[("1","Problem Identification & Research","2 Weeks"),("2","Literature Survey","2 Weeks"),("3","System Design","3 Weeks"),("4","Implementation (Core Modules)","4 Weeks"),("5","Plugin Development","3 Weeks"),("6","Testing & Evaluation","2 Weeks"),("7","Documentation & Review","2 Weeks")])
ip("Figure 4.5 \u2013 Gantt Chart"); pb()

# === CHAPTER 5 ===
h1("Chapter 5: Implementation Results and Discussions")
h2("5.1 Data Collection / Acquisition")
bd('Dione operates on dynamic, real-time data: user input (natural language commands), system context (active window, file system), local storage (documents, PDFs), external APIs (Gmail, Calendar), and ChromaDB (embeddings of past interactions).')
h2("5.2 Data Preprocessing")
bd('Techniques: Text Normalization, Tokenization, Context Augmentation via RAG, Schema Structuring to JSON, and Validation Preprocessing using Pydantic.')
h2("5.3 Experimental Setup")
bd('Hardware: i5/Ryzen 5, 16-32 GB RAM, RTX 3060/Apple Silicon, SSD. Software: Python 3.10+, Ollama/Llama.cpp, ChromaDB, SQLite, Pydantic. Tests: Single-step (fetch file), Multi-step (retrieve, summarize, email), Context-based (use history).')
ip("Figure 5.1 \u2013 CLI Screenshot: File Creation Task"); ip("Figure 5.2 \u2013 CLI Screenshot: WhatsApp Message Task"); ip("Figure 5.3 \u2013 CLI Screenshot: Gmail Read/Send Task")
h2("5.4 Evaluation Metrics")
bd('Task Success Rate (TSR): percentage of completed tasks. Time-to-Action (TTA): input to first action time. Token Generation Speed: tokens/sec. Privacy Metric: zero unauthorized data leakage. Safety Metric: 100% prevention of unsafe commands without approval.')
h2("5.5 Gap Analysis")
tbl(["Feature","Existing Systems","Dione"],[("Execution","Limited","Full task execution"),("Privacy","Low (cloud)","High (local)"),("Security","Vulnerable","Controlled"),("Latency","High (network)","Low (local)"),("Autonomy","Partial","High (multi-step)"),("Scalability","Limited","Modular plugin-based")]); pb()

# === CHAPTER 6 ===
h1("Chapter 6: Plan of Action for Next Semester")
h2("6.1 Work Done Till Date")
bd('Completed: Problem identification (Actuation Gap analysis), Literature Survey (Toolformer, ReAct, RAG), System Design (architecture, block/modular diagrams, ReAct loop), Requirement Analysis, Initial Implementation (Ollama/Llama.cpp setup, basic pipeline, initial plugin structure), and Documentation.')
h2("6.2 Plan of Action for Project II")
bd('Core System Development: Complete orchestration engine, improve reasoning-to-action pipeline, enhance LLM-plugin integration. Plugin Expansion: Email automation, calendar scheduling, file system management, API integrations. Security Enhancement: Strengthen guardrails, permission-based execution, confirmation layer. Performance Optimization: Quantization, reduced latency, improved memory management. UI Development: Web-based GUI. Testing: AgentBench benchmarks, TSR/Latency/Safety evaluation. Final Integration and Deployment.'); pb()

# === CHAPTER 7 ===
h1("Chapter 7: Conclusion")
bd('The development of Dione represents a significant advancement in addressing the Actuation Gap. While modern LLMs have demonstrated remarkable reasoning capabilities, they remain constrained by their inability to execute real-world actions. This project bridges that gap by introducing an intelligent orchestration layer enabling seamless interaction between AI reasoning and system-level execution.')
bd('The core contribution lies in its local-first architecture, which ensures all reasoning, context handling, and decision-making occur entirely on the user\u2019s device, guaranteeing data sovereignty, enhanced privacy, and reduced security risks.')
bd('The system integrates the ReAct loop, RAG, and a modular plugin-based execution framework. These components work cohesively to handle complex, multi-step workflows dynamically. The ReAct loop ensures iterative reasoning and correction, the context manager enhances decision-making, and the plugin architecture ensures scalability and extensibility.'); pb()

# === CHAPTER 8 ===
h1("Chapter 8: References")
for r in ['[1] Schick, T., et al. (2023). "Toolformer: Language Models Can Teach Themselves to Use Tools." arXiv:2302.04761.','[2] Wei, J., et al. (2022). "Chain-of-Thought Prompting Elicits Reasoning in LLMs." NeurIPS.','[3] Nakajima, Y. (2023). "BabyAGI: Task-driven Autonomous Agent." GitHub.','[4] OpenAI. (2023). "GPT-4 Technical Report."','[5] Chase, H. (2022). "LangChain: Building applications with LLMs."','[6] Yao, S., et al. (2023). "ReAct: Synergizing Reasoning and Acting in LLMs." arXiv:2210.03629.','[7] Touvron, H., et al. (2023). "Llama 2: Open Foundation and Fine-Tuned Chat Models." Meta AI.','[8] Jiang, A.Q., et al. (2023). "Mistral 7B." arXiv:2310.06825.','[9] Lewis, P., et al. (2020). "RAG for Knowledge-Intensive NLP Tasks." NeurIPS.','[10] Dettmers, T., et al. (2023). "QLoRA: Efficient Finetuning of Quantized LLMs." arXiv:2305.14314.','[11] Vaswani, A., et al. (2017). "Attention Is All You Need." NeurIPS.','[12] Park, J.S., et al. (2023). "Generative Agents." arXiv:2304.03442.','[13] Mialon, G., et al. (2023). "Augmented Language Models: A Survey." arXiv:2302.07842.','[14] Significant-Gravitas. (2023). "AutoGPT." GitHub.','[15] Kandpal, N., et al. (2023). "LLMs Struggle to Learn Long-Tail Knowledge." arXiv:2211.08411.']: bd(r)
pb()

# === APPENDIX ===
h1("Appendix")
h2("A. List of Figures")
tbl(["Figure No.","Heading","Page No."],[("4.1","Block Diagram",""),("4.2","Modular Diagram",""),("4.3","Use Case Diagram",""),("4.4","System Flowchart",""),("4.5","Gantt Chart",""),("5.1","CLI: File Creation",""),("5.2","CLI: WhatsApp",""),("5.3","CLI: Gmail","")])
doc.add_paragraph()
h2("B. List of Tables")
tbl(["Table No.","Heading","Page No."],[("4.1","Project Schedule",""),("5.1","Gap Analysis",""),("CO","Course Outcomes","")])
doc.add_paragraph()
h2("C. Paper Publications")
bd("[Acceptance mail, Published paper, Plagiarism report, and Certificate \u2013 To be attached.]")

# === SAVE ===
out = 'docs/Dione_Capstone_Report_Final.docx'
doc.save(out)
print(f"Report saved to: {out}")
