import docx
import os

doc = docx.Document(r"c:\Users\Jass\Documents\CAPSTONE\PROJECT - I\docs\Dione_Report_v2.docx")
print("Total paragraphs:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._element.xml for r in p.runs)
    if has_img:
        prev_t = doc.paragraphs[i-1].text[:60] if i>0 else 'START'
        print(f"P[{i}]: [IMAGE] Text: '{p.text[:60]}' | Prev: '{prev_t}'")
